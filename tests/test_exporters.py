"""Output policy tests; processed values must remain unchanged."""
from contextlib import ExitStack
from datetime import date, datetime, time, timezone
from decimal import Decimal
import csv
import json
from pathlib import Path
import tempfile
import tracemalloc
import unittest
from uuid import UUID
from unittest.mock import patch

from etl.engine import execute
from etl.exporters import OutputWriter, RejectedWriter
from etl.models import FieldError, RowResult, SourceRow
from etl.serialization import pipeline_from_dict, pipeline_to_dict
from etl.spec import ConfigError, validate


def result(values, errors=()):
    return RowResult(SourceRow(7, values, 9, 11), values, values, errors)


class ExporterTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def export(self, values, *, version=2, kind="csv", **options):
        with ExitStack() as stack:
            writer = OutputWriter(self.root, values.keys(), {"kind": kind, **options}, stack, version=version)
            stack.callback(writer.finish)
            writer.write_result(result(values))
        return writer.path

    def csv_rows(self, path, delimiter=";", encoding="utf-8-sig"):
        with path.open(encoding=encoding, newline="") as handle:
            return list(csv.reader(handle, delimiter=delimiter, strict=True))

    def test_v2_csv_scalar_formats_and_exact_text_roundtrip(self):
        values = {"null": None, "empty": "", "code": "003", "greek": "Αθήνα",
                  "space": " \t ", "delimiter": "a;b", "quotes": 'a"b',
                  "multiline": "\r\nπρώτη\nδεύτερη\rτέλος", "decimal": Decimal("12345678901234567890.1234500"),
                  "integer": 9007199254740993, "true": True, "false": False,
                  "date": date(2024, 2, 29), "datetime": datetime(2024, 2, 29, 12, 34, 56, 123000, timezone.utc),
                  "time": time(12, 34, 56), "bytes": b"\x00\xff", "uuid": UUID(int=3)}
        path = self.export(values)
        self.assertTrue(path.read_bytes().startswith(b"\xef\xbb\xbf"))
        rows = self.csv_rows(path)
        self.assertEqual(rows[0], list(values))
        self.assertEqual(rows[1], ["\\N", "", "003", "Αθήνα", " \t ", "a;b", 'a"b',
                                  "\r\nπρώτη\nδεύτερη\rτέλος", "12345678901234567890.1234500", "9007199254740993",
                                  "true", "false", "2024-02-29", "2024-02-29T12:34:56.123000+00:00", "12:34:56",
                                  "base64:AP8=", "00000000-0000-0000-0000-000000000003"])
        self.assertIsNone(values["null"])
        self.assertEqual(values["code"], "003")

    def test_utf8_delimiter_quoting_and_record_terminator(self):
        path = self.export({"κεφαλίδα": 'Α, "Β"\nΓ'}, encoding="utf-8", delimiter=",")
        self.assertEqual(path.read_bytes(), 'κεφαλίδα\r\n"Α, ""Β""\nΓ"\r\n'.encode("utf-8"))
        self.assertEqual(self.csv_rows(path, ",", "utf-8")[1], ['Α, "Β"\nΓ'])

    def test_explicit_null_tokens_and_explicit_lossy_empty_policy(self):
        for token in ("", "NULL", "\\N", 'missing;"value"\n'):
            with self.subTest(token=token):
                path = self.export({"null": None, "empty": ""}, null_value=token)
                self.assertEqual(self.csv_rows(path)[1], [token, ""])

    def test_null_token_collisions_fail_for_strings_and_formatted_numbers(self):
        for token, value, policy in (("\\N", "\\N", "preserve"), ("0", 0, "preserve"),
                                     ("NULL", "NULL", "apostrophe"), ("=NULL", "=NULL", "apostrophe")):
            with self.subTest(token=token), self.assertRaisesRegex(ConfigError, "collides"):
                self.export({"value": value}, null_value=token, formula_policy=policy)

    def test_formula_policies_are_explicit_and_do_not_mutate_results(self):
        values = {str(n): value for n, value in enumerate(("=1+1", "+003", "-003", " @SUM(A1)",
                  "\tfoo", "\rfoo", "\nfoo", "a\nb", "'=literal", -12, Decimal("-1.2500")))}
        original = result(values)
        preserved = self.csv_rows(self.export(values))[1]
        self.assertEqual(preserved, [str(v) for v in values.values()])
        protected = self.csv_rows(self.export(values, formula_policy="apostrophe"))[1]
        self.assertEqual(protected, ["'=1+1", "'+003", "'-003", "' @SUM(A1)", "'\tfoo", "'\rfoo", "'\nfoo",
                                     "a\nb", "'=literal", "-12", "-1.2500"])
        self.assertEqual(dict(original.converted_values), values)

    def test_formula_headers_follow_policy(self):
        self.assertEqual(self.csv_rows(self.export({"=header": "x"}, formula_policy="apostrophe"))[0], ["'=header"])
        self.assertEqual(self.csv_rows(self.export({"=header": "x"}))[0], ["=header"])

    def test_v1_csv_golden_output_preserved(self):
        path = self.export({"null": None, "empty": "", "code": "003", "formula": "=1+1",
                            "number": Decimal("-123.4500"), "boolean": False}, version=1)
        self.assertEqual(path.read_bytes(), b"\xef\xbb\xbfnull;empty;code;formula;number;boolean\r\n;;003;'=1+1;-123.4500;false\r\n")

    def test_v2_xlsx_uses_text_cells_with_distinct_null_and_empty(self):
        from openpyxl import load_workbook
        values = {"null": None, "empty": "", "code": "003", "decimal": Decimal("12345678901234567890.1234500"),
                  "integer": 9007199254740993, "formula": "=1+1", "boolean": True, "date": date(2024, 2, 29),
                  "datetime": datetime(2024, 2, 29, 1, 2, 3), "greek": "Αθήνα", "multiline": "α\r\nβ\nγ\rδ", "space": "   "}
        path = self.export(values, kind="xlsx", formula_policy="apostrophe")
        workbook = load_workbook(path, read_only=True)
        try:
            cells = list(workbook.active.rows)[1]
            self.assertEqual([c.value for c in cells], ["\\N", None, "003", "12345678901234567890.1234500",
                             "9007199254740993", "=1+1", "true", "2024-02-29", "2024-02-29T01:02:03", "Αθήνα", "α\r\nβ\nγ\rδ", "   "])
            self.assertTrue(all(c.data_type == "s" for n, c in enumerate(cells) if n != 1))
            self.assertEqual(cells[1].data_type, "inlineStr")
        finally:
            workbook.close()

    def test_v1_xlsx_null_empty_and_formula_behavior_preserved(self):
        from openpyxl import load_workbook
        path = self.export({"null": None, "empty": "", "formula": "=1+1", "code": "003"}, version=1, kind="xlsx")
        workbook = load_workbook(path, read_only=True)
        try:
            cells = list(workbook.active.rows)[1]
            self.assertEqual([c.value for c in cells], [None, None, "=1+1", "003"])
            self.assertEqual(cells[2].data_type, "s")
        finally:
            workbook.close()

    def test_xlsx_rejects_overlong_text_instead_of_truncating(self):
        with self.assertRaisesRegex(ConfigError, "32,767"):
            self.export({"value": "x" * 32768}, kind="xlsx")

    def test_xlsx_backend_and_unrepresentable_values_fail_explicitly(self):
        from openpyxl.utils.exceptions import IllegalCharacterError
        with patch("openpyxl.LXML", False), self.assertRaisesRegex(ConfigError, "lxml backend"):
            self.export({"value": "α\r\nβ"}, kind="xlsx")
        with self.assertRaises(IllegalCharacterError):
            self.export({"value": "a\x00b"}, kind="xlsx")
        with self.assertRaisesRegex(ConfigError, "collides"):
            self.export({"value": "\\N"}, kind="xlsx")

    def test_decimal_exponents_and_nonfinite_or_unknown_types(self):
        self.assertEqual(self.csv_rows(self.export({"d": Decimal("1.2300E+20")}))[1], ["1.2300E+20"])
        for value in (Decimal("NaN"), float("inf"), object()):
            with self.subTest(value=type(value).__name__), self.assertRaises(ConfigError):
                # Exercise the exporter boundary directly; these values cannot
                # reach a valid RowResult through normal strict processing.
                with ExitStack() as stack:
                    writer = OutputWriter(self.root, ["v"], {"kind": "csv"}, stack, version=2)
                    writer.write({"v": value})

    def test_v2_execute_honors_output_policy_without_changing_preview_values(self):
        (self.root / "input.csv").write_bytes(b"id\n1\n")
        spec = {"version": 2, "name": "policy integration", "source": {"kind": "csv", "path": "input.csv"},
                "columns": [{"name": "null", "literal": None}, {"name": "empty", "literal": ""},
                            {"name": "decimal", "literal": "123.4500", "type": "decimal"},
                            {"name": "formula", "literal": "=1+1"}],
                "destination": {"kind": "csv", "encoding": "utf-8", "delimiter": ",", "null_value": "NULL", "formula_policy": "apostrophe"}}
        preview = execute(spec, self.root, limit=1)
        full = execute(spec, self.root, self.root / "out")
        self.assertEqual(preview["sample"], full["sample"])
        self.assertEqual(dict(full["sample"][0].converted_values), {"null": None, "empty": "", "decimal": Decimal("123.4500"), "formula": "=1+1"})
        rows = self.csv_rows(Path(full["directory"]) / "valid.csv", ",", "utf-8")
        self.assertEqual(rows[1], ["NULL", "", "123.4500", "'=1+1"])

    def test_xlsx_sheet_rollover_repeats_headers(self):
        from openpyxl import load_workbook
        with ExitStack() as stack:
            writer = OutputWriter(self.root, ["code"], {"kind": "xlsx"}, stack, version=2)
            writer.XLSX_DATA_ROWS = 2
            stack.callback(writer.finish)
            for value in ("001", "002", "003"):
                writer.write_result(result({"code": value}))
        workbook = load_workbook(writer.path, read_only=True)
        try:
            self.assertEqual([list(sheet.values) for sheet in workbook], [[("code",), ("001",), ("002",)], [("code",), ("003",)]])
        finally:
            workbook.close()

    def test_rejections_use_separate_stages_and_native_json_tags(self):
        source = SourceRow(7, {"amount": Decimal("123.4500"), "bad": " 003 ", "empty": "", "null": None, "memo": 'α;"β"\nγ'}, 9, 11)
        rejected = RowResult(source, {"amount": Decimal("123.4500"), "bad": "003"}, {"amount": Decimal("123.4500")},
                             (FieldError("bad", "invalid_type", "expected integer", "conversion"),))
        with ExitStack() as stack:
            writer = RejectedWriter(self.root, {"kind": "csv", "delimiter": ",", "encoding": "utf-8"}, stack, version=2)
            writer.write_result(rejected)
        rows = self.csv_rows(self.root / "rejected.csv", ",", "utf-8")
        self.assertEqual(rows[0], list(RejectedWriter.HEADERS))
        record, original, stages, errors = rows[1]
        self.assertEqual(record, "7")
        original, stages, errors = map(json.loads, (original, stages, errors))
        self.assertIsNone(original["null"])
        self.assertEqual(original["empty"], "")
        self.assertEqual(original["memo"], 'α;"β"\nγ')
        self.assertEqual(stages["source_position"], {"line_start": 9, "line_end": 11})
        self.assertEqual(stages["transformed_values"]["bad"], "003")
        self.assertNotIn("bad", stages["converted_values"])
        self.assertEqual(stages["converted_values"]["amount"], {"$type": "decimal", "value": "123.4500"})
        self.assertEqual(errors["bad"], [{"field": "bad", "code": "invalid_type", "stage": "conversion", "message": "expected integer"}])

    def test_v1_rejection_layout_and_mixed_projection_remain_compatible(self):
        row = RowResult(SourceRow(1, {"bad": " x ", "empty": ""}), {"bad": "x", "empty": ""}, {"empty": None},
                        (FieldError("bad", "invalid_type", "bad integer", "conversion"),))
        with ExitStack() as stack:
            writer = RejectedWriter(self.root, {"kind": "csv", "delimiter": ","}, stack, version=1)
            writer.write_result(row)
        rows = self.csv_rows(self.root / "rejected.csv")
        self.assertEqual(rows[0], list(RejectedWriter.HEADERS))
        self.assertEqual(json.loads(rows[1][2]), {"bad": "x", "empty": None})
        self.assertEqual(json.loads(rows[1][3]), {"bad": ["bad integer"]})

    def test_valid_and_rejected_exporters_enforce_classification(self):
        with ExitStack() as stack:
            writer = OutputWriter(self.root, ["v"], {"kind": "csv"}, stack, version=2)
            rejected = RejectedWriter(self.root, {"kind": "csv"}, stack, version=2)
            with self.assertRaisesRegex(ConfigError, "rejected RowResult"):
                writer.write_result(result({"v": "x"}, (FieldError("v", "required", "missing", "required"),)))
            with self.assertRaisesRegex(ConfigError, "valid RowResult"):
                rejected.write_result(result({"v": "x"}))

    def test_large_exports_keep_bounded_memory_and_all_rows(self):
        from openpyxl import load_workbook  # Import before measuring allocations.
        for kind, count in (("csv", 20000), ("xlsx", 5000)):
            with self.subTest(kind=kind):
                tracemalloc.start()
                try:
                    with ExitStack() as stack:
                        writer = OutputWriter(self.root, ["id", "memo"], {"kind": kind}, stack, version=2)
                        stack.callback(writer.finish)
                        for n in range(count):
                            writer.write_result(result({"id": n, "memo": "α" * 200 + "\nβ"}))
                        self.assertEqual(writer.count, count)
                        if kind == "xlsx":
                            self.assertTrue(writer.workbook.write_only)
                    _, peak = tracemalloc.get_traced_memory()
                finally:
                    tracemalloc.stop()
                self.assertLess(peak, 8 * 1024 * 1024)
                if kind == "csv":
                    with writer.path.open(encoding="utf-8-sig", newline="") as handle:
                        self.assertEqual(sum(1 for _ in csv.reader(handle, delimiter=";", strict=True)), count + 1)
                else:
                    workbook = load_workbook(writer.path, read_only=True)
                    try:
                        self.assertEqual(sum(1 for _ in workbook.active.rows), count + 1)
                    finally:
                        workbook.close()

    def test_policy_definition_roundtrip_and_v1_rejection_of_new_options(self):
        definition = {"version": 2, "name": "export policy", "source": {"kind": "csv", "path": "input.csv"},
                      "columns": [{"name": "code", "source": "code"}],
                      "destination": {"kind": "csv", "encoding": "utf-8", "delimiter": ",", "null_value": "NULL", "formula_policy": "apostrophe"}}
        model = pipeline_from_dict(definition)
        self.assertEqual(pipeline_to_dict(model), definition)
        with self.assertRaises(TypeError):
            model.destination["null_value"] = ""
        definition["version"] = 1
        with self.assertRaisesRegex(ConfigError, "Unknown destination"):
            validate(definition)
        definition["version"] = 2
        for key, value in (("encoding", "latin1"), ("null_value", None), ("formula_policy", "guess")):
            invalid = {**definition, "destination": {**definition["destination"], key: value}}
            with self.subTest(key=key), self.assertRaises(ConfigError):
                validate(invalid)
