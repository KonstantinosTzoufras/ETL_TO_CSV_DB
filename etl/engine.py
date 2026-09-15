"""Pure mapping and validation with bounded previews and streamed exports."""
import csv
import json
import math
import re
import uuid
from contextlib import ExitStack
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from itertools import islice
from pathlib import Path

from .sources import open_source
from .spec import ConfigError, require, validate


def text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def transform(value, column):
    for operation in column.get("transforms", []):
        if operation == "empty_to_null":
            if value == "":
                value = None
        elif value is not None:
            value = {"trim": str.strip, "upper": str.upper, "lower": str.lower}[operation](text(value))
    return value


def convert(value, kind):
    if value is None or value == "":
        return None
    value = text(value)
    if kind == "string":
        return value
    if kind == "int":
        if not re.fullmatch(r"-?(0|[1-9][0-9]*)", value) or not -(2**63) <= int(value) < 2**63:
            raise ValueError("expected a 64-bit integer without leading zeros")
        return int(value)
    if kind in {"decimal", "float"}:
        if not re.fullmatch(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?", value):
            raise ValueError("expected a number using a decimal point")
        number = Decimal(value) if kind == "decimal" else float(value)
        if not (number.is_finite() if isinstance(number, Decimal) else math.isfinite(number)):
            raise ValueError("number must be finite")
        return number
    if kind == "bool":
        if value.lower() not in {"0", "1", "true", "false"}:
            raise ValueError("expected 0, 1, true or false")
        return value.lower() in {"1", "true"}
    if kind == "date":
        if not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", value):
            raise ValueError("expected YYYY-MM-DD")
        return date.fromisoformat(value)
    if kind == "datetime":
        if not re.match(r"[0-9]{4}-[0-9]{2}-[0-9]{2}[T ]", value):
            raise ValueError("expected an ISO date and time")
        return datetime.fromisoformat(value)
    raise ConfigError(f"Unsupported type: {kind}")


def lookup_sets(spec, root):
    result = {}
    for column in spec["columns"]:
        if "lookup" not in column:
            continue
        lookup = column["lookup"]
        allowed = set()
        with open_source(lookup["source"], root) as (headers, rows):
            require(lookup["column"] in headers, f"Unknown lookup column: {lookup['column']}")
            for number, row in enumerate(rows):
                require(number < 100000, "Lookup exceeds 100,000 rows; narrow the lookup source")
                value = transform(row[lookup["column"]], column)
                if value is not None and value != "":
                    try:
                        allowed.add(convert(value, column.get("type", "string")))
                    except (ValueError, InvalidOperation, OverflowError):
                        raise ConfigError(f"Lookup contains an invalid {column.get('type', 'string')} value") from None
        result[column["name"]] = allowed
    return result


def map_row(row, columns, lookups):
    output, errors = {}, {}
    for column in columns:
        name = column["name"]
        value = transform(row[column["source"]] if "source" in column else column["literal"], column)
        reasons = []
        if column.get("required") and (value is None or not text(value).strip()):
            reasons.append("required value is missing")
        if "max_length" in column and len(text(value)) > column["max_length"]:
            reasons.append(f"exceeds {column['max_length']} characters")
        try:
            converted = convert(value, column.get("type", "string"))
            if name in lookups and converted is not None and converted not in lookups[name]:
                reasons.append("value is absent from lookup")
            output[name] = converted
        except (ValueError, InvalidOperation, OverflowError) as error:
            output[name] = value
            reasons.append(str(error))
        if reasons:
            errors[name] = reasons
    return output, errors


def excel_safe(value):
    """Escape formula-like text in CSVs intended to be opened in Excel."""
    if not isinstance(value, str):
        return text(value)
    value = text(value)
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")) else value


class OutputWriter:
    def __init__(self, directory, names, destination, stack):
        self.kind = destination["kind"]
        self.path = directory / ("valid." + self.kind)
        self.names = names
        self.workbook = None
        self.count = 0
        if self.kind == "csv":
            handle = stack.enter_context(self.path.open("w", encoding="utf-8-sig", newline=""))
            self.writer = csv.writer(handle, delimiter=destination.get("delimiter", ";"))
            self.writer.writerow([excel_safe(n) for n in names])
        else:
            try:
                from openpyxl import Workbook
            except ImportError:
                raise ConfigError("XLSX export needs openpyxl. Install requirements.txt.") from None
            self.workbook = Workbook(write_only=True)
            self.sheet = self.workbook.create_sheet("Data")
            self.append_xlsx(names)

    def append_xlsx(self, values):
        from openpyxl.cell import WriteOnlyCell
        cells = []
        for value in values:
            value = text(value)
            require(len(value) <= 32767, "XLSX cell exceeds 32,767 characters; use CSV")
            cell = WriteOnlyCell(self.sheet, value=value)
            cell.data_type = "s"  # Preserve IDs and decimals; never create formulas.
            cells.append(cell)
        self.sheet.append(cells)

    def write(self, row):
        values = [row[name] for name in self.names]
        if self.kind == "csv":
            self.writer.writerow([excel_safe(value) for value in values])
        else:
            if self.count and self.count % 1048575 == 0:
                self.sheet = self.workbook.create_sheet()
                self.append_xlsx(self.names)
            self.append_xlsx(values)
        self.count += 1

    def finish(self):
        if self.workbook:
            self.workbook.save(self.path)
            self.workbook.close()


def execute(spec, root, output_root=None, limit=None, progress=None):
    validate(spec)
    require(limit is None or type(limit) is int and 1 <= limit <= 1000, "Preview limit must be 1–1000")
    require(limit is not None or output_root is not None, "Full execution needs an output directory")
    names = [column["name"] for column in spec["columns"]]
    lookups = lookup_sets(spec, root)
    report = {"processed": 0, "valid": 0, "invalid": 0, "sample": [], "preview": limit is not None}
    with ExitStack() as stack:
        headers, rows = stack.enter_context(open_source(spec["source"], root))
        for column in spec["columns"]:
            require("source" not in column or column["source"] in headers, f"Unknown source column: {column.get('source')}")
        writer = rejected = None
        if output_root is not None and limit is None:
            directory = Path(output_root) / uuid.uuid4().hex
            directory.mkdir(parents=True)
            report["directory"] = str(directory)
            writer = OutputWriter(directory, names, spec["destination"], stack)
            # Finalize write-only XLSX streams even if a later input row fails.
            # Failed runs never expose their partial files for download.
            stack.callback(writer.finish)
            rejected = csv.writer(stack.enter_context((directory / "rejected.csv").open("w", encoding="utf-8-sig", newline="")), delimiter=";")
            rejected.writerow(["record_number", "source_json", "mapped_json", "errors_json"])
        for number, row in enumerate(islice(rows, limit) if limit is not None else rows, 1):
            mapped, errors = map_row(row, spec["columns"], lookups)
            report["processed"] += 1
            report["invalid" if errors else "valid"] += 1
            if len(report["sample"]) < (limit or 20):
                report["sample"].append({"record": number, "values": {k: text(v) for k, v in mapped.items()}, "errors": errors})
            if errors and rejected:
                rejected.writerow([number, json.dumps(row, ensure_ascii=False, default=text), json.dumps(mapped, ensure_ascii=False, default=text), json.dumps(errors, ensure_ascii=False)])
            elif writer and not errors:
                writer.write(mapped)
            if progress and number % 1000 == 0:
                progress({k: report[k] for k in ("processed", "valid", "invalid")})
    return report
