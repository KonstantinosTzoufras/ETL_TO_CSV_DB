"""Streaming output boundaries. Never transform or infer processed values."""
import base64
import csv
import json
import math
from datetime import date, datetime, time
from decimal import Decimal
from uuid import UUID

from .conversions import text
from .serialization import json_default, row_result_to_dict
from .spec import ConfigError, destination_spec, require


def legacy_projection(result):
    """Historical presentation shape, retained for v1 callers and rejections."""
    values = {name: result.converted_values.get(name, value)
              for name, value in result.transformed_values.items()}
    errors = {}
    for error in result.errors:
        errors.setdefault(error.field, []).append(error.message)
    return values, errors


def excel_safe(value):
    """Historical CSV policy: apostrophe-prefix formula-like strings only."""
    if not isinstance(value, str):
        return text(value)
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")) else value


def scalar_text(value):
    """V2 non-null output formats, independent of CSV/XLSX cell quoting."""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, Decimal):
        require(value.is_finite(), "Cannot export a non-finite Decimal")
        return str(value)  # Preserve scale/exponent directly; never use float.
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        require(math.isfinite(value), "Cannot export a non-finite float")
        return str(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, bytes):
        return "base64:" + base64.b64encode(value).decode("ascii")
    raise ConfigError(f"Unsupported export value type: {type(value).__name__}")


class OutputWriter:
    """write_result(valid RowResult), finish(); memory bounded to one output row.

    write(mapping) and the version=1 default retain the original public API.
    The caller owns the ExitStack and registers finish() after construction.
    """
    XLSX_DATA_ROWS = 1048575  # One header row per sheet.

    def __init__(self, directory, names, destination, stack, *, version=1):
        destination_spec(destination, version)
        self.version = version
        self.kind = destination["kind"]
        self.path = directory / ("valid." + self.kind)
        self.names = tuple(names)
        self.encoding = destination.get("encoding", "utf-8-sig")
        self.null_value = destination.get("null_value", "" if version == 1 else "\\N")
        self.formula_policy = destination.get("formula_policy", "apostrophe" if version == 1 else "preserve")
        self.workbook = None
        self.count = 0
        self.finished = False
        if self.kind == "csv":
            handle = stack.enter_context(self.path.open("w", encoding=self.encoding, newline=""))
            self.writer = csv.writer(handle, delimiter=destination.get("delimiter", ";"),
                                     quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
            self.writer.writerow([self._protect(name) for name in self.names])
        else:
            try:
                from openpyxl import Workbook, LXML
            except ImportError:
                raise ConfigError("XLSX export needs openpyxl. Install requirements.txt.") from None
            require(version == 1 or LXML, "V2 XLSX requires the lxml backend to preserve line endings. Install requirements.txt and restart, or use CSV.")
            self.workbook = Workbook(write_only=True)
            self.sheet = self.workbook.create_sheet("Data")
            self.append_xlsx(self.names)

    def _protect(self, value):
        return excel_safe(value) if self.kind == "csv" and self.formula_policy == "apostrophe" else value

    def format_value(self, value):
        if self.version == 1:
            return excel_safe(value) if self.kind == "csv" else text(value)
        null = self._protect(self.null_value)
        if value is None:
            return null
        rendered = scalar_text(value)
        # Native numbers remain numeric representations (e.g. -12), not
        # formula-like strings. String values and bytes get the selected policy.
        if isinstance(value, (str, bytes)):
            rendered = self._protect(rendered)
        if self.null_value != "":
            require(rendered != null, "Export value collides with null_value; choose a different NULL token")
        return rendered

    def append_xlsx(self, values):
        from openpyxl.cell import WriteOnlyCell
        cells = []
        for value in values:
            require(len(value) <= 32767, "XLSX cell exceeds 32,767 characters; use CSV")
            cell = WriteOnlyCell(self.sheet, value=value)
            cell.data_type = "s"  # Includes identifiers, decimals, dates and '=' text.
            cells.append(cell)
        self.sheet.append(cells)

    def write_result(self, result):
        require(result.valid, "Valid-row exporter received a rejected RowResult")
        self.write(result.converted_values)

    def write(self, row):
        require(not self.finished, "Exporter is already finished")
        values = [self.format_value(row[name]) for name in self.names]
        if self.kind == "csv":
            self.writer.writerow(values)
        else:
            if self.count and self.count % self.XLSX_DATA_ROWS == 0:
                self.sheet = self.workbook.create_sheet()
                self.append_xlsx(self.names)
            self.append_xlsx(values)
        self.count += 1

    def finish(self):
        if self.finished:
            return
        self.finished = True
        if self.workbook:
            try:
                self.workbook.save(self.path)
            finally:
                self.workbook.close()


class RejectedWriter:
    """Structured v2 rejection JSON inside the existing four-column CSV layout."""
    HEADERS = ("record_number", "source_json", "mapped_json", "errors_json")

    def __init__(self, directory, destination, stack, *, version):
        destination_spec(destination, version)
        self.version = version
        handle = stack.enter_context((directory / "rejected.csv").open(
            "w", encoding=destination.get("encoding", "utf-8-sig"), newline=""))
        self.writer = csv.writer(handle, delimiter=";" if version == 1 else destination.get("delimiter", ";"),
                                 quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
        self.writer.writerow(self.HEADERS)

    def write_result(self, result):
        require(not result.valid, "Rejected-row exporter received a valid RowResult")
        if self.version == 1:
            mapped, errors = legacy_projection(result)
            payload = [json.dumps(dict(result.original_values), ensure_ascii=False, default=text),
                       json.dumps(mapped, ensure_ascii=False, default=text), json.dumps(errors, ensure_ascii=False)]
        else:
            diagnostic = row_result_to_dict(result)
            stages = {key: diagnostic[key] for key in ("source_position", "transformed_values", "converted_values")}
            payload = [json.dumps(value, ensure_ascii=False, default=json_default)
                       for value in (diagnostic["original_values"], stages, diagnostic["errors"])]
        self.writer.writerow([result.source.number, *payload])
